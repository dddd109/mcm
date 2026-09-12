# -*- coding: utf-8 -*-
"""2026 CUMCM A 题「药材烘干」论文配图（国赛标准，mcm-plot 层）。
数据源: D:\\CUMCM2026Problems\\A题\\最终代码与结果\\outputs
输出:   ......\\figs_论文配图\\  (PNG + SVG，逐张过 guard 质检)
"""
import os, sys, json, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Ellipse, FancyArrowPatch, Circle, FancyArrow

SKILL = r"C:\Users\AD\.config\opencode\skills\mcm-plot"
sys.path.insert(0, SKILL)
from style import new_axes, SAVE, COLORS, PALETTE, apply_style, ensure_cn_font
from captions import put_caption, caption_rule_check, annotate_help
from guard import check, assert_clean

B = r"D:\CUMCM2026Problems\A题\最终代码与结果\outputs"
OUT = r"D:\CUMCM2026Problems\A题\最终代码与结果\figs_论文配图"
os.makedirs(OUT, exist_ok=True)
ensure_cn_font(verbose=True)
CMAP_T = "YlOrRd"
CMAP_C = "YlGnBu"


def load_xlsx(path, sheets):
    wb = openpyxl.load_workbook(path, data_only=True)
    res = {}
    for sh in sheets:
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        labels = list(header[1:])
        radii = np.array([float(x) if isinstance(x, (int, float)) else np.nan for x in labels])
        data = rows[1:]
        times = np.array([float(r[0]) for r in data])
        M = np.array([[float(v) if isinstance(v, (int, float)) else np.nan for v in r[1:]] for r in data], dtype=float)
        res[sh] = (times, labels, radii, M)
    wb.close()
    return res


def emit(fig, name, fig_no, title, detail=None, result=None, vs_patch=True):
    caption_rule_check(fig, fig.axes[0] if fig.axes else None, name=fig_no, verbose=True)
    res = check(fig, name=fig_no, vs_ticks=True, vs_patch=vs_patch)
    try:
        assert_clean(res, name=fig_no)
    except Exception as e:
        print("  [guard] 未过:", e)
    put_caption(fig, fig_no, title, detail=detail, result=result)
    SAVE(fig, os.path.join(OUT, name))


# ============================================================ 图1 模型示意
def fig01():
    fig, axs = new_axes(7.6, 0.44, 1, 2)
    for ax in axs:
        ax.grid(False)
        ax.set_axis_off()
        ax.set_xlim(0, 10); ax.set_ylim(0, 6)
    ax = axs[0]
    ax.add_patch(Rectangle((1.4, 2.0), 7.0, 2.0, facecolor="#EDEDED", edgecolor="#222222", lw=1.0, zorder=1))
    ax.add_patch(Ellipse((1.4, 3.0), 1.1, 2.0, facecolor="#F5F5F5", edgecolor="#222222", lw=1.0, zorder=2))
    ax.add_patch(Ellipse((8.4, 3.0), 1.1, 2.0, facecolor="#EDEDED", edgecolor="#222222", lw=1.0, zorder=2))
    ax.plot([1.4, 8.4], [3.0, 3.0], ls=(0, (5, 4)), color="#888888", lw=0.9)
    ax.annotate("", xy=(8.9, 1.2), xytext=(1.4, 1.2),
                arrowprops=dict(arrowstyle="<->", color="#222222", lw=0.9))
    ax.text(5.0, 0.85, "长度 $L=25\\,$cm", ha="center", fontsize=9)
    ax.annotate("", xy=(1.4, 4.9), xytext=(4.0, 4.9), arrowprops=dict(arrowstyle="<->", color="#222222", lw=0.9))
    ax.text(2.7, 5.1, "半径 $R_0=2\\,$cm", ha="center", fontsize=9)
    for x0 in (3.0, 4.4, 5.8):
        ax.annotate("", xy=(x0, 4.0), xytext=(x0, 4.7),
                    arrowprops=dict(arrowstyle="-|>", color=COLORS["red"], lw=1.0))
    ax.text(2.2, 4.35, "$h_T,\\,h_m$", color=COLORS["red"], fontsize=9)
    ax.text(6.6, 4.55, "$T_\\infty,\\,C_\\infty$", color="#333333", fontsize=9)
    ax.text(5.0, 2.75, "药材（多孔介质）", ha="center", fontsize=9, color="#333333")
    ax.text(0.1, 5.6, "(a) 圆柱药材与对流边界", fontsize=9.5)

    ax = axs[1]
    ax.add_patch(Circle((5.0, 3.0), 1.9, facecolor="#EDEDED", edgecolor="#222222", lw=1.0, zorder=1))
    ax.add_patch(Circle((5.0, 3.0), 0.05, facecolor="#222222"))
    ax.annotate("", xy=(6.9, 3.0), xytext=(5.0, 3.0), arrowprops=dict(arrowstyle="-|>", color="#222222", lw=1.0))
    ax.text(5.9, 3.18, "$r$", fontsize=10)
    ax.text(6.6, 2.7, "$R(t)$", fontsize=9)
    for ang in range(0, 360, 45):
        a = np.deg2rad(ang)
        x0, y0 = 5.0 + 1.9 * np.cos(a), 3.0 + 1.9 * np.sin(a)
        x1, y1 = 5.0 + 2.6 * np.cos(a), 3.0 + 2.6 * np.sin(a)
        ax.annotate("", xy=(x1, y1), xytext=(x0, y0),
                    arrowprops=dict(arrowstyle="-|>", color=COLORS["blue"], lw=1.0))
    ax.text(5.0, 0.55, "表面：$-k\\,\\partial T/\\partial r=h_T(T_\\infty-T_s)$，$-D\\,\\partial C/\\partial r=h_m(C_s-C_\\infty)$",
            ha="center", fontsize=8.3, color="#333333")
    ax.text(0.1, 5.6, "(b) 横截面与第三类边界", fontsize=9.5)
    emit(fig, "fig01_模型示意图.png", "图 1", "圆柱药材二维轴对称传热-传质模型与边界条件示意",
         detail="侧面与两端面均为对流边界；$r$ 为到中轴距离",
         result="热扩散时标 ~40 min、湿扩散时标 ~天量级，两尺度分离", vs_patch=False)


# ============================================================ 图2 问题一 时空场
def fig02():
    d = load_xlsx(os.path.join(B, "problem1", "result1.xlsx"), ["温度", "水分浓度"])
    fig, axs = new_axes(8.0, 0.42, 1, 2)
    fig.subplots_adjust(bottom=0.24, top=0.84, wspace=0.34)
    for i, (ax, sh, cmap, lab) in enumerate(
            ((axs[0], "温度", CMAP_T, "温度 $T$ / $^\\circ$C"),
             (axs[1], "水分浓度", CMAP_C, "含水率 $C$ / (kg$\\cdot$kg$^{-1}$)"))):
        t, labels, radii, M = d[sh]
        tt = t / 60.0
        pc = ax.pcolormesh(tt, radii, M.T, cmap=cmap, shading="auto")
        pc.set_rasterized(True)
        cb = fig.colorbar(pc, ax=ax, pad=0.02)
        cb.set_label(lab, fontsize=9)
        ax.set_xlabel("时间 $t$ / min")
        if i == 0:
            ax.set_ylabel("到药材中心距离 $r$ / cm")
    axs[0].set_title("(a) 温度场", fontsize=10)
    axs[1].set_title("(b) 含水率场", fontsize=10)
    emit(fig, "fig02_问题一_时空演化.png", "图 2", "问题一预热段温度场与含水率场的时空演化",
         detail="$0\\sim1800\\,$s，附录 2 常物性，附件 1 环境",
         result="$1800\\,$s 时温度中心 33.58$\\,^\\circ$C、表面 36.79$\\,^\\circ$C；含水率仅表面下降、中心几乎不变")


# ============================================================ 图3 问题一 径向剖面
def fig03():
    d = load_xlsx(os.path.join(B, "problem1", "result1.xlsx"), ["温度", "水分浓度"])
    times = [300, 600, 900, 1200, 1800]
    cols = ["#4C72B0", "#55A868", "#B5BD61", "#DD8452", "#C44E52"]
    fig, axs = new_axes(8.0, 0.40, 1, 2)
    fig.subplots_adjust(bottom=0.26, top=0.84, wspace=0.34)
    t, labels, radii, MT = d["温度"]
    _, _, _, MC = d["水分浓度"]
    for c, ts in zip(cols, times):
        i = int(np.argmin(np.abs(t - ts)))
        axs[0].plot(radii, MT[i], color=c, label=f"$t={ts}\\,$s")
        axs[1].plot(radii, MC[i], color=c, label=f"$t={ts}\\,$s")
    for ax, yl, ttl in ((axs[0], "温度 $T$ / $^\\circ$C", "(a) 温度径向剖面"),
                        (axs[1], "含水率 $C$ / (kg$\\cdot$kg$^{-1}$)", "(b) 含水率径向剖面")):
        ax.set_xlabel("到药材中心距离 $r$ / cm")
        ax.set_ylabel(yl)
        ax.set_title(ttl, fontsize=10)
        ax.legend(fontsize=8, loc="best")
    emit(fig, "fig03_问题一_径向剖面.png", "图 3", "问题一不同时刻的温度与含水率径向剖面",
         detail="由中心向外，表面处梯度最陡",
         result="温度剖面逐渐抬升趋近环境；含水率仍保持外低内高的核状分布")


# ============================================================ 图4 问题二 时空场
def fig04():
    z = np.load(os.path.join(B, "problem2", "fields_N100_dt1.npz"))
    t = z["times_s"]; rc = z["r_centers_m"] * 100.0
    T = z["T_C"]; C = z["C_kg_kg"]
    sl = slice(0, None, 10); rsl = slice(0, None, 1)
    fig, axs = new_axes(8.0, 0.42, 1, 2)
    fig.subplots_adjust(bottom=0.24, top=0.84, wspace=0.34)
    for i, (ax, Z, cmap, lab) in enumerate(
            ((axs[0], T, CMAP_T, "温度 $T$ / $^\\circ$C"),
             (axs[1], C, CMAP_C, "含水率 $C$ / (kg$\\cdot$kg$^{-1}$)"))):
        pc = ax.pcolormesh(t[sl] / 3600.0, rc[rsl], Z[sl, rsl].T, cmap=cmap, shading="auto")
        pc.set_rasterized(True)
        cb = fig.colorbar(pc, ax=ax, pad=0.02); cb.set_label(lab, fontsize=9)
        ax.set_xlabel("时间 $t$ / h")
        if i == 0:
            ax.set_ylabel("到药材中心距离 $r$ / cm")
    axs[0].set_title("(a) 温度场", fontsize=10)
    axs[1].set_title("(b) 含水率场", fontsize=10)
    emit(fig, "fig04_问题二_时空演化.png", "图 4", "问题二耦合变物性下温度场与含水率场的时空演化",
         detail="$0\\sim3\\,$h，附录 3 变物性，$t>4\\,$h 环境保护末值",
         result="3 h 时温度中心 49.85$\\,^\\circ$C、表面 49.97$\\,^\\circ$C；含水率表面降至约 1.01、中心约 1.77")


# ============================================================ 图5 问题二 时序
def fig05():
    z = np.load(os.path.join(B, "problem2", "fields_N100_dt1.npz"))
    t = z["times_s"] / 3600.0
    pos = z["positions_m"] * 100.0
    ST = z["sampled_T_C"]; SC = z["sampled_C_kg_kg"]
    ambT = z["ambient_T_C"]; ambC = z["ambient_C_kg_kg"]
    i0 = 0; imid = int(np.argmin(np.abs(pos - 1.0))); isf = len(pos) - 1
    fig, axs = new_axes(8.0, 0.40, 1, 2)
    fig.subplots_adjust(bottom=0.26, top=0.84, wspace=0.34)
    axs[0].plot(t, ambT, color="#999999", ls=(0, (5, 3)), label="环境 $T_\\infty$")
    axs[0].plot(t, ST[:, imid], color=COLORS["orange"], label="中间 $r=1$ cm")
    axs[0].plot(t, ST[:, i0], color=COLORS["blue"], label="中心 $r=0$")
    axs[0].plot(t, ST[:, isf], color=COLORS["red"], label="表面 $r=2$ cm")
    axs[0].set_ylabel("温度 $T$ / $^\\circ$C"); axs[0].set_xlabel("时间 $t$ / h")
    axs[0].set_title("(a) 温度随时间", fontsize=10); axs[0].legend(fontsize=8)
    axs[1].plot(t, ambC, color="#999999", ls=(0, (5, 3)), label="环境 $C_\\infty$")
    axs[1].plot(t, SC[:, imid], color=COLORS["orange"], label="中间 $r=1$ cm")
    axs[1].plot(t, SC[:, i0], color=COLORS["blue"], label="中心 $r=0$")
    axs[1].plot(t, SC[:, isf], color=COLORS["red"], label="表面 $r=2$ cm")
    axs[1].set_ylabel("含水率 $C$ / (kg$\\cdot$kg$^{-1}$)"); axs[1].set_xlabel("时间 $t$ / h")
    axs[1].set_title("(b) 含水率随时间", fontsize=10); axs[1].legend(fontsize=8)
    emit(fig, "fig05_问题二_特征点时序.png", "图 5", "问题二中心、中间与表面的温度及含水率时序",
         detail="环境取附件 1，$t>4\\,$h 后维持 $T_\\infty=50\\,^\\circ$C、$C_\\infty=0.05$",
         result="温度先快速上升后趋缓并趋近环境；含水率下降由表及里，中心显著滞后于表面")


# ============================================================ 图6 问题三 时空场
def fig06():
    d = load_xlsx(os.path.join(B, "problem3", "result3.xlsx"), ["Sheet1"])
    t, labels, radii, M = d["Sheet1"]
    fig, ax = new_axes(6.4, 0.62)
    pc = ax.pcolormesh(t / 3600.0, radii, M.T, cmap=CMAP_C, shading="auto")
    pc.set_rasterized(True)
    cb = fig.colorbar(pc, ax=ax, pad=0.02); cb.set_label("含水率 $C$ / (kg$\\cdot$kg$^{-1}$)", fontsize=9)
    td = 207360 / 3600.0
    ax.axvline(td, color=COLORS["red"], lw=1.3, ls="--")
    ax.text(td - 1.0, 1.85, f"$t_d={td:.1f}\\,$h", color=COLORS["red"], ha="right", fontsize=9)
    ax.set_xlabel("时间 $t$ / h"); ax.set_ylabel("到药材中心距离 $r$ / cm")
    emit(fig, "fig06_问题三_含水率时空演化.png", "图 6", "问题三含水率场时空演化直至处处达标",
         detail="每 60 s 采样，判据 $\\max_{r}C\\leq0.15\\,$kg/kg",
         result=f"$t_d={td:.4f}\\,$h（{207360}\\,s）时各处含水率低于 0.15")


# ============================================================ 图7 问题三 达标曲线
def fig07():
    d = load_xlsx(os.path.join(B, "problem3", "result3.xlsx"), ["Sheet1"])
    t, labels, radii, M = d["Sheet1"]
    tg = t / 3600.0
    c0 = M[:, 0]; cs = M[:, -1]; cmax = np.nanmax(M, axis=1)
    imid = int(np.argmin(np.abs(radii - 1.0))); cmid = M[:, imid]
    fig, ax = new_axes(6.4, 0.58)
    ax.plot(tg, cmax, color=COLORS["red"], label="处处最大 $\\max_{r} C$")
    ax.plot(tg, c0, color=COLORS["blue"], label="中心 $r=0$")
    ax.plot(tg, cmid, color=COLORS["orange"], label="中间 $r=1$ cm")
    ax.plot(tg, cs, color=COLORS["green"], label="表面 $r=2$ cm")
    ax.axhline(0.15, color="#333333", lw=1.0, ls=":")
    ax.text(1.0, 0.19, "达标线 $C=0.15$", fontsize=8.5, color="#333333")
    td = 207360 / 3600.0
    ax.axvline(td, color="#333333", lw=1.0, ls="--")
    ax.plot([td], [0.15], marker="o", ms=4, color=COLORS["red"])
    ax.annotate(f"$t_d={td:.2f}\\,$h", xy=(td, 0.15), xytext=(td - 14, 0.55),
                fontsize=9, color=COLORS["red"], arrowprops=dict(arrowstyle="-|>", color=COLORS["red"], lw=0.8))
    ax.set_xlabel("时间 $t$ / h"); ax.set_ylabel("含水率 $C$ / (kg$\\cdot$kg$^{-1}$)")
    ax.set_ylim(0, 2.6); ax.legend(fontsize=8.5)
    emit(fig, "fig07_问题三_达标过程.png", "图 7", "问题三各处含水率随时间下降与达标时刻",
         detail="判据为重建圆心、全部控制体中心与真实表面含水率的最大值",
         result=f"$t_d={td:.4f}\\,$h；根因：中心为最慢控制点，始终为 $\\max_{{r}} C$ 的所在")


# ============================================================ 图8 问题四
def fig08():
    d = load_xlsx(os.path.join(B, "problem4", "result4.xlsx"), ["Sheet1"])
    t, labels, radii, M = d["Sheet1"]
    tg = t / 3600.0
    # 半径 R(t)（附件2 线性插值）
    wb = openpyxl.load_workbook(os.path.join(B, "problem1", "result1.xlsx"), data_only=True)
    wb.close()
    # 附件2: 0..259200 s @1800
    Rcm = None
    for cand in [r"D:\CUMCM2026Problems\A题\附件\附件2.xlsx"]:
        if os.path.exists(cand):
            ws = openpyxl.load_workbook(cand, data_only=True).active
            rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
            tt = np.array([float(r[0]) for r in rows]); RR = np.array([float(r[1]) for r in rows])
            Rcm = np.interp(t, tt, RR)
    c0 = M[:, 0]
    csurf = np.array([M[i, -1] for i in range(M.shape[0])], dtype=float)
    cmax = np.nanmax(M, axis=1)
    fig, axs = new_axes(8.0, 0.40, 1, 2)
    fig.subplots_adjust(bottom=0.26, top=0.84, wspace=0.34)
    ax = axs[0]
    if Rcm is not None:
        ax.plot(tg, Rcm, color=COLORS["purple"], lw=1.8)
        ax.set_ylabel("药材半径 $R(t)$ / cm"); ax.set_xlabel("时间 $t$ / h")
        ax.set_title("(a) 半径收缩过程", fontsize=10)
        ax.set_ylim(0, 2.2)
    ax = axs[1]
    ax.plot(tg, cmax, color=COLORS["red"], label="处处最大 $\\max_{r} C$")
    ax.plot(tg, c0, color=COLORS["blue"], label="中心 $r=0$")
    ax.plot(tg, csurf, color=COLORS["green"], label="药材表面")
    ax.axhline(0.15, color="#333333", lw=1.0, ls=":")
    td4 = 184012 / 3600.0
    ax.axvline(td4, color="#333333", lw=1.0, ls="--")
    ax.text(td4 - 0.8, 1.9, f"$t_{{d4}}={td4:.2f}\\,$h", color=COLORS["red"], ha="right", fontsize=9)
    ax.set_ylabel("含水率 $C$ / (kg$\\cdot$kg$^{-1}$)"); ax.set_xlabel("时间 $t$ / h")
    ax.set_title("(b) 含水率与达标时刻", fontsize=10); ax.legend(fontsize=8)
    emit(fig, "fig08_问题四_收缩与达标.png", "图 8", "问题四尺寸收缩下的半径变化与含水率达标",
         detail="材料坐标 $\\xi=r/R(t)$，附录 4 物性，$R$ 由附件 2 给定",
         result=f"半径由 2.0 收缩至 1.20 cm；$t_{{d4}}={td4:.4f}\\,$h，较问题三缩短约 11.26%")


# ============================================================ 图9 问题三 vs 问题四
def fig09():
    d3 = load_xlsx(os.path.join(B, "problem3", "result3.xlsx"), ["Sheet1"])
    d4 = load_xlsx(os.path.join(B, "problem4", "result4.xlsx"), ["Sheet1"])
    t3, _, _, M3 = d3["Sheet1"]; t4, _, _, M4 = d4["Sheet1"]
    tg3 = t3 / 3600.0; tg4 = t4 / 3600.0
    c3 = np.nanmax(M3, axis=1); c4 = np.nanmax(M4, axis=1)
    fig, ax = new_axes(6.4, 0.58)
    ax.plot(tg3, c3, color=COLORS["blue"], label="问题三（$R$ 不变，附录 3）")
    ax.plot(tg4, c4, color=COLORS["red"], label="问题四（$R$ 收缩，附录 4）")
    ax.axhline(0.15, color="#333333", lw=1.0, ls=":")
    ax.axvline(207360 / 3600, color=COLORS["blue"], lw=0.9, ls="--")
    ax.axvline(184012 / 3600, color=COLORS["red"], lw=0.9, ls="--")
    ax.plot([207360 / 3600], [0.15], "o", color=COLORS["blue"], ms=4)
    ax.plot([184012 / 3600], [0.15], "o", color=COLORS["red"], ms=4)
    ax.text(207360 / 3600 - 1, 0.55, "$t_d=57.60\\,$h", color=COLORS["blue"], ha="right", fontsize=8.5)
    ax.text(184012 / 3600 - 1, 0.35, "$t_{d4}=51.11\\,$h", color=COLORS["red"], ha="right", fontsize=8.5)
    ax.set_xlabel("时间 $t$ / h"); ax.set_ylabel("处处最大含水率 $\\max_{r} C$ / (kg$\\cdot$kg$^{-1}$)")
    ax.set_ylim(0, 2.6); ax.legend(fontsize=8.5)
    emit(fig, "fig09_问题三四_对比.png", "图 9", "问题三与问题四的处处最大含水率下降过程对比",
         detail="问题四同时含半径收缩与附录 4 物性变化",
         result="收缩缩短扩散路径（加速）与附录 4 扩散系数变小（减速）竞争，净效应使达标时间提前 6.49 h")


# ============================================================ 图10 收敛验证
def fig10():
    cm = json.load(open(os.path.join(B, "problem1", "convergence_metrics.json"), encoding="utf-8"))
    p3 = json.load(open(os.path.join(B, "problem3", "validation_summary.json"), encoding="utf-8"))
    p4 = json.load(open(os.path.join(B, "problem4", "validation_summary.json"), encoding="utf-8"))
    fig, axs = new_axes(8.0, 0.40, 1, 2)
    fig.subplots_adjust(bottom=0.26, top=0.84, wspace=0.34)
    ax = axs[0]
    labels = ["$N$:50→100", "$N$:100→200"]
    Tmax = [cm["space_50_100"]["T"]["max_all"], cm["space_100_200"]["T"]["max_all"]]
    Cmax = [cm["space_50_100"]["C"]["max_all"], cm["space_100_200"]["C"]["max_all"]]
    x = np.arange(2); w = 0.36
    ax.bar(x - w / 2, Tmax, w, color=COLORS["blue"], label="温度 $T$ / $^\\circ$C")
    ax.bar(x + w / 2, Cmax, w, color=COLORS["orange"], label="含水率 $C$ / (kg/kg)")
    ax.set_xticks(x); ax.set_xticklabels(labels)
    ax.set_ylabel("相邻网格最大偏差")
    ax.set_title("(a) 空间加密收敛", fontsize=10); ax.legend(fontsize=8)
    ax = axs[1]
    Ns = [50, 100, 200]
    td3 = {r["N"]: r["t_dry_h"] for r in p3["runs"] if r["dt_s"] == 1.0}
    td4 = {r["N"]: r["t_dry_h"] for r in p4["runs"] if r["dt_s"] == 1.0}
    ax.plot(Ns, [td3[n] for n in Ns], "o-", color=COLORS["blue"], label="问题三 $t_d$")
    ax.plot(Ns, [td4[n] for n in Ns], "s-", color=COLORS["red"], label="问题四 $t_{d4}$")
    ax.set_xscale("log"); ax.set_xticks(Ns); ax.set_xticklabels(["50", "100", "200"])
    ax.set_xlabel("径向网格数 $N$"); ax.set_ylabel("烘干时长 / h")
    ax.set_title("(b) 达标时长随网格收敛", fontsize=10); ax.legend(fontsize=8)
    emit(fig, "fig10_收敛验证.png", "图 10", "网格加密下的数值收敛验证",
         detail="空间加密使相邻网格最大偏差下降；达标时长随 $N$ 增趋于稳定",
         result="问题三/四达标时长在 $N=100\\to200$ 仅差 294 s / 58 s，满足网格无关性")


if __name__ == "__main__":
    for fn in [fig01, fig02, fig03, fig04, fig05, fig06, fig07, fig08, fig09, fig10]:
        print("### ", fn.__name__)
        fn()
    print("ALL DONE ->", OUT)
